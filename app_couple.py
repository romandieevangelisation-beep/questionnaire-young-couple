import streamlit as st
import pandas as pd
import plotly.graph_objects as go
from datetime import datetime
from fpdf import FPDF
import os
import tempfile
import re
# Nouvelle librairie pour lire les fichiers Word
try:
    from docx import Document
except ImportError:
    st.error("Erreur : La librairie 'python-docx' est manquante. Ajoutez-la dans requirements.txt")

# --- 0. CONFIGURATION ---
st.set_page_config(page_title="Alliance & Schémas - Expert", layout="wide", page_icon="✝️")
DB_FILE = "reponses_couple_master.csv"

# --- 1. BIBLIOTHÈQUE D'EXPERTISE (CLINIQUE, THÉOLOGIQUE & PASTORALE) ---
SCHEMA_LIBRARY = {
    "ca": { # ED dans votre fichier
        "nom": "Carence Affective",
        "clinique": "Sentiment profond que ses besoins de sécurité, d'affection ou d'empathie ne seront jamais comblés. Le patient se sent invisible et émotionnellement privé.",
        "theologie": "Le Syndrome de l'Orphelin. C'est la croyance mensongère que Dieu est avare ou distant. Le cœur cherche à combler une soif infinie avec des citernes crevassées (Jérémie 2:13).",
        "couple": "Le 'Trou Noir'. Le patient attend (souvent en silence) que l'autre devine ses besoins. Quand l'autre échoue, il punit par le froid ou la colère. Le conjoint se sent impuissant et finit par se désinvestir.",
        "pratique": "1. Apprendre la demande directe : 'J'ai besoin d'un câlin' (Vulnérabilité). 2. Deuil des parents idéaux. 3. Recevoir l'amour de Dieu par la contemplation.",
        "verset": "« D'un amour éternel je t'ai aimé, c'est pourquoi je t'ai conservé ma bonté. » (Jérémie 31:3)"
    },
    "ab": { # AB
        "nom": "Abandon / Instabilité",
        "clinique": "Peur envahissante que les proches meurent, partent ou trouvent mieux ailleurs. Hypersensibilité à toute forme de séparation.",
        "theologie": "Idolâtrie de la Sécurité Relationnelle. On demande à l'humain la fidélité absolue que seul Dieu possède. C'est un doute profond sur l'Alliance et la Providence.",
        "couple": "Le 'Velcro'. Jalousie, possessivité, 'tests' de l'amour. Étouffe le conjoint, ce qui provoque paradoxalement le recul de l'autre (prophétie auto-réalisatrice).",
        "pratique": "1. Exercice de 'Solitude habitée' avec Dieu. 2. Arrêter les SMS de vérification. 3. S'ancrer dans la permanence de Christ.",
        "verset": "« Je ne te délaisserai point, et je ne t'abandonnerai point. » (Hébreux 13:5)"
    },
    "ma": { # MA
        "nom": "Méfiance / Abus",
        "clinique": "Attente que les autres vont nuire, manipuler ou trahir. Hypervigilance. Difficulté extrême à faire confiance et à s'abandonner.",
        "theologie": "La blessure de Caïn. Le monde est vu comme une jungle hostile sans la protection de Dieu. La 'Crainte de l'homme' remplace la confiance en Dieu comme Bouclier (Psaume 3).",
        "couple": "La 'Forteresse'. Le patient interprète les erreurs involontaires du conjoint comme des attaques malveillantes. Il garde des secrets et ne baisse jamais la garde.",
        "pratique": "1. Arrêter la 'lecture de pensée'. 2. Oser une confiance vérifiable sur des petites choses. 3. Pardonner les offenses passées pour ne pas punir le conjoint actuel.",
        "verset": "« L'amour parfait bannit la crainte. » (1 Jean 4:18)"
    },
    "is": { # SI
        "nom": "Isolement Social",
        "clinique": "Sentiment d'être différent, de ne pas appartenir au groupe, d'être exclu ou inintéressant socialement.",
        "theologie": "Refus de la Communion Fraternelle. Soit par Orgueil (se croire unique/supérieur), soit par Honte (se croire indigne). C'est une coupure avec le Corps du Christ.",
        "couple": "Le 'Couple Ermite'. Le patient isole le couple, refusant les amis ou la famille. Il demande au conjoint d'être son 'seul univers', une charge trop lourde pour un humain.",
        "pratique": "1. S'engager dans un service concret (laver les pieds des autres brise l'isolement). 2. Pratiquer l'hospitalité.",
        "verset": "« Dieu donne une famille à ceux qui étaient abandonnés. » (Psaume 68:6)"
    },
    "im": { # DS
        "nom": "Imperfection / Honte",
        "clinique": "Sentiment d'être intérieurement défectueux, mauvais, sans valeur. Peur qu'on découvre notre 'vraie' nature et qu'on soit rejeté.",
        "theologie": "La Honte d'Adam (Genèse 3). Difficulté à accepter la Justification par la Foi. On veut 'payer' pour son péché ou le cacher, au lieu de le confesser et recevoir la Grâce.",
        "couple": "La 'Défensive'. Très susceptible à la critique. Soit il s'écrase (confirme sa honte), soit il contre-attaque violemment pour cacher sa vulnérabilité.",
        "pratique": "1. Transparence radicale avec Dieu et le conjoint. 2. Distinguer comportement (ce que je fais) et identité (qui je suis en Christ).",
        "verset": "« Il n'y a donc maintenant aucune condamnation pour ceux qui sont en Jésus-Christ. » (Romains 8:1)"
    },
    "ed": { # FA
        "nom": "Échec",
        "clinique": "Croyance d'être moins capable que les autres, inepte, voué à l'échec professionnel ou personnel. Comparaison constante.",
        "theologie": "L'Idole de la Réussite. La valeur personnelle est attachée à la performance (œuvres) et non à l'adoption filiale. C'est une forme d'orgueil inversé.",
        "couple": "L'Enfant incompétent'. Se repose entièrement sur le conjoint pour les décisions 'sérieuses', créant un déséquilibre parent/enfant qui tue l'admiration mutuelle.",
        "pratique": "1. Redéfinir le succès biblique : la fidélité, pas le résultat. 2. Prendre une responsabilité concrète par semaine.",
        "verset": "« C'est bien, bon et fidèle serviteur ; tu as été fidèle en peu de chose. » (Matthieu 25:21)"
    },
    "da": { # DI
        "nom": "Dépendance / Incompétence",
        "clinique": "Incapacité à gérer le quotidien sans l'aide d'autrui. Se sent comme un enfant dans un monde d'adultes.",
        "theologie": "Refus de la responsabilité d'Intendant. C'est refuser de porter sa propre charge (Galates 6:5) en cherchant un 'sauveur humain' à la place de l'Esprit Saint.",
        "couple": "Le 'Fardeau'. L'un porte tout, l'autre suit. Au début c'est flatteur pour le 'fort', à la fin c'est épuisant et cela tue le désir sexuel.",
        "pratique": "1. Prise de décision autonome quotidienne (même petite). 2. Le conjoint doit arrêter de valider chaque choix.",
        "verset": "« Car ce n'est pas un esprit de timidité que Dieu nous a donné, mais un esprit de force, d'amour et de sagesse. » (2 Timothée 1:7)"
    },
    "vu": { # VU
        "nom": "Vulnérabilité au Danger",
        "clinique": "Peur constante et irrationnelle qu'une catastrophe (médicale, financière, criminelle) va survenir.",
        "theologie": "Manque de foi en la Providence. L'inquiétude est un 'athéisme pratique' : on vit comme si Dieu ne contrôlait pas l'univers. Esclavage de la peur de la mort.",
        "couple": "La 'Prison de Sécurité'. Empêche le couple de prendre des risques, de voyager, d'investir. Le conjoint est utilisé comme une police d'assurance.",
        "pratique": "1. Transformer chaque inquiétude en prière. 2. Exposition progressive aux situations redoutées. 3. Lâcher prise.",
        "verset": "« Ne vous inquiétez de rien; mais en toute chose faites connaître vos besoins à Dieu. » (Philippiens 4:6)"
    },
    "fu": { # EU
        "nom": "Fusion / Personnalité Atrophiée",
        "clinique": "Trop impliqué émotionnellement avec les parents ou le conjoint. Pas d'identité propre, sentiments en miroir.",
        "theologie": "Idolâtrie relationnelle. Ne pas avoir 'quitté son père et sa mère'. L'autre prend la place de Dieu comme source de vie.",
        "couple": "Le 'Siamois'. Aucune intimité personnelle. Si le conjoint est triste, le patient est dévasté. Étouffement mutuel et perte de désir.",
        "pratique": "1. Développer des hobbies séparés. 2. Apprendre à dire 'Je' au lieu de 'Nous'. 3. Couper le cordon émotionnel.",
        "verset": "« C'est pourquoi l'homme quittera son père et sa mère, et s'attachera à sa femme. » (Genèse 2:24)"
    },
    "ass": { # SB
        "nom": "Assujettissement",
        "clinique": "Se soumettre au contrôle des autres par peur de la colère ou du rejet. Refoulement de ses propres besoins et émotions.",
        "theologie": "Crainte de l'homme (Proverbes 29:25). On sert la créature plutôt que le Créateur. On achète une fausse paix au prix de la Vérité.",
        "couple": "La 'Cocotte-Minute'. Dit 'oui' à tout mais accumule une rancœur secrète. Finit par exploser ou devenir passif-agressif (sabotage inconscient).",
        "pratique": "1. Apprendre l'assertivité chrétienne : dire la vérité dans l'amour. 2. Dire 'Non' est spirituel quand c'est juste.",
        "verset": "« Si je plaisais encore aux hommes, je ne serais pas serviteur de Christ. » (Galates 1:10)"
    },
    "ss": { # SS
        "nom": "Sacrifice de Soi",
        "clinique": "Se concentrer excessivement sur les besoins des autres au détriment des siens (Syndrome du Sauveur).",
        "theologie": "Orgueil caché sous l'humilité. Penser qu'on est le 'Messie' de l'autre. Confusion entre amour (vouloir le bien) et sacrifice névrotique (acheter l'amour).",
        "couple": "Le 'Martyr'. Donne tout, s'épuise, puis culpabilise le conjoint ('Après tout ce que j'ai fait pour toi !'). Crée une dette relationnelle.",
        "pratique": "1. Recevoir est aussi important que donner. 2. Sabbat (repos) obligatoire. 3. Examiner ses motivations : est-ce par amour ou par culpabilité ?",
        "verset": "« Tu aimeras ton prochain comme toi-même. » (Marc 12:31)"
    },
    "ie": { # EI
        "nom": "Inhibition Émotionnelle",
        "clinique": "Verrouillage des émotions, de la spontanéité. Paraît froid, rationnel, robotique par peur de perdre le contrôle.",
        "theologie": "Stoïcisme vs Incarnation. Jésus a pleuré, a exulté. Refus de la condition humaine sensible créée par Dieu. Cœur de pierre.",
        "couple": "Le 'Mur de Glace'. Le conjoint se sent seul et non-aimé. Pas de joie, pas de rire. La relation devient fonctionnelle et morte.",
        "pratique": "1. Journal des émotions (Lire les Psaumes). 2. Oser dire 'Je suis triste' ou 'Je suis joyeux'. 3. Partager sans rationaliser.",
        "verset": "« Je vous donnerai un cœur nouveau... j'ôterai de votre corps le cœur de pierre. » (Ézéchiel 36:26)"
    },
    "is_std": { # US
        "nom": "Exigences Élevées",
        "clinique": "Standards de performance inatteignables. Perfectionnisme rigide. Jamais satisfait, critique envers soi et les autres.",
        "theologie": "Pharisaïsme. On place la Loi au-dessus de la Grâce. On juge les autres comme on pense que Dieu nous juge (par les œuvres).",
        "couple": "Le 'Juge'. Rien n'est jamais assez bien (ménage, éducation, salaire). Le conjoint vit sous pression constante et finit par démissionner.",
        "pratique": "1. Célébrer 'l'assez bien'. 2. Accepter l'imperfection comme une réalité de la Chute. 3. Pratiquer la gratitude pour ce qui est fait.",
        "verset": "« Car c'est par la grâce que vous êtes sauvés... cela ne vient pas de vous, c'est le don de Dieu. » (Éphésiens 2:8)"
    },
    "dt": { # ET
        "nom": "Droits Personnels / Grandeur",
        "clinique": "Se sentir supérieur, spécial, au-dessus des lois. Manque d'empathie pour les besoins des autres. Colère si frustré.",
        "theologie": "Orgueil pur (Péché originel : 'Vous serez comme des dieux'). Refus de la condition de serviteur. C'est l'anti-thèse de l'esprit de Christ.",
        "couple": "Le 'Tyran'. Le conjoint est un objet utilitaire. Colère narcissique si les besoins ne sont pas satisfaits immédiatement.",
        "pratique": "1. Service humble et anonyme. 2. Empathie cognitive : 'Comment se sent l'autre ?'. 3. Se rappeler que nous sommes poussière.",
        "verset": "« Ne faites rien par esprit de parti ou par vaine gloire, mais que l'humilité vous fasse regarder les autres comme étant au-dessus de vous-mêmes. » (Philippiens 2:3)"
    },
    "ci": { # IS (User file conflict solved: IS=Control here)
        "nom": "Contrôle de Soi Insuffisant",
        "clinique": "Impulsivité, intolérance à la frustration, addictions, procrastination. Difficulté à différer la gratification.",
        "theologie": "Esclavage des pulsions (le 'Ventre' comme dieu). Manque de 'Fruit de l'Esprit' (Maîtrise de soi). Vie dictée par le plaisir immédiat.",
        "couple": "L'Enfant Capricieux'. On ne peut pas compter sur lui. Dépenses impulsives, paroles blessantes, infidélité possible par manque de frein.",
        "pratique": "1. Tolérance à la frustration (Jeûne). 2. Mettre un temps d'arrêt entre l'impulsion et l'action. 3. Redevabilité.",
        "verset": "« Comme une ville forcée et sans murailles, ainsi est l'homme qui n'est pas maître de lui-même. » (Proverbes 25:28)"
    },
    "rc": { # AS
        "nom": "Recherche d'Approbation",
        "clinique": "Recherche excessive de l'attention et de la validation. Estime de soi dépendante du regard des autres.",
        "theologie": "Idolâtrie de la Gloire humaine. On préfère la louange des hommes à celle de Dieu. On est un 'caméléon' sans colonne vertébrale.",
        "couple": "L'Acteur'. Change de personnalité selon qui est là. Le conjoint ne sait plus qui il aime vraiment. Jalousie sociale.",
        "pratique": "1. Vivre 'Coram Deo' (Devant la face de Dieu seul). 2. Faire des choses 'en secret' pour Dieu. 3. Authenticité radicale.",
        "verset": "« Comment pouvez-vous croire, vous qui tirez votre gloire les uns des autres, et qui ne cherchez point la gloire qui vient de Dieu seul? » (Jean 5:44)"
    },
    "neg": { # NP
        "nom": "Négativisme / Pessimisme",
        "clinique": "Focalisation sur les aspects négatifs (douleur, mort, perte, conflit). Inquiétude chronique et plaintes.",
        "theologie": "Ingratitude et Manque d'Espérance. Cécité face à la Grâce commune et à la bonté de Dieu. Cœur amer.",
        "couple": "Le 'Rabat-joie'. Éteint l'enthousiasme du conjoint. 'Ça ne marchera pas'. Ambiance lourde et déprimante à la maison.",
        "pratique": "1. Discipline de la Gratitude (Noter 3 grâces par jour). 2. Louange intentionnelle. 3. Se forcer à voir le verre à moitié plein.",
        "verset": "« Rendez grâces en toutes choses, car c'est à votre égard la volonté de Dieu en Jésus-Christ. » (1 Thessaloniciens 5:18)"
    },
    "pu": { # PU
        "nom": "Punition",
        "clinique": "Intolérance, critique, tendance à punir durement les erreurs (soi et les autres). Difficulté à pardonner.",
        "theologie": "Légalisme et Refus de la Miséricorde. Dieu est vu comme un Père fouettard. Oubli de sa propre dette infinie envers Dieu.",
        "couple": "Le 'Bourreau'. Rancunier. Sort les vieux dossiers lors des disputes. Le conjoint a peur de faire une erreur. Climat de peur.",
        "pratique": "1. Méditer la parabole du débiteur impitoyable. 2. Pardonner c'est renoncer à se venger. 3. Grâce vs Justice.",
        "verset": "« Soyez bons les uns envers les autres, compatissants, vous pardonnant réciproquement, comme Dieu vous a pardonné en Christ. » (Éphésiens 4:32)"
    }
}
SCHEMAS_ORDER = list(SCHEMA_LIBRARY.keys())

# --- 2. ENGINE : MAPPING & QUESTIONS ---

def get_schema_map_ordered():
    """Mapping EXACT basé sur votre fichier Word (Reponses_Peter Pan)"""
    m = []
    m.extend(['ca'] * 9)   # Q1-9 Carence
    m.extend(['ab'] * 17)  # Q10-26 Abandon
    m.extend(['ma'] * 17)  # Q27-43 Méfiance
    m.extend(['is'] * 10)  # Q44-53 Isolement
    m.extend(['im'] * 15)  # Q54-68 Imperfection
    m.extend(['ed'] * 9)   # Q69-77 Echec
    m.extend(['da'] * 15)  # Q78-92 Dépendance
    m.extend(['vu'] * 12)  # Q93-104 Vulnérabilité
    m.extend(['fu'] * 11)  # Q105-115 Fusion
    m.extend(['ass'] * 10) # Q116-125 Assujettissement
    m.extend(['ss'] * 17)  # Q126-142 Sacrifice
    m.extend(['ie'] * 9)   # Q143-151 Inhibition
    m.extend(['is_std'] * 16) # Q152-167 Exigences
    m.extend(['dt'] * 11)  # Q168-178 Droits
    m.extend(['ci'] * 15)  # Q179-193 Contrôle Insuffisant
    m.extend(['rc'] * 14)  # Q194-207 Recherche Approbation
    m.extend(['neg'] * 15) # Q208-222 Négativisme
    m.extend(['pu'] * 10)  # Q223-232 Punition
    return m

def generate_web_questions():
    """Génère les questions placeholders pour le web"""
    q = {}
    m = get_schema_map_ordered()
    for idx, sc in enumerate(m):
        q[idx+1] = {"text": f"Question {idx+1} ({SCHEMA_LIBRARY[sc]['nom']})", "schema": sc}
    return q

ALL_QUESTIONS = generate_web_questions()

# --- NOUVEAU : FONCTION DE LECTURE DOCX ---
def extract_text_from_file(uploaded_file):
    """Lit indifféremment un .txt ou un .docx"""
    if uploaded_file.name.endswith('.docx'):
        doc = Document(uploaded_file)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return '\n'.join(full_text)
    else:
        # Fichier texte classique
        return uploaded_file.getvalue().decode("utf-8", "ignore")

def parse_imported_text(text_content):
    """Analyse le texte brut pour trouver les [x/6]"""
    matches = re.findall(r"\[(\d)/6\]", text_content)
    if not matches: return None, "Aucune note [x/6] trouvée."
    
    scores = [int(x) for x in matches]
    mapping = get_schema_map_ordered()
    limit = min(len(scores), len(mapping))
    
    sums = {s:0 for s in SCHEMAS_ORDER}
    cnts = {s:0 for s in SCHEMAS_ORDER}
    
    for i in range(limit):
        sch = mapping[i]
        sums[sch] += scores[i]
        cnts[sch] += 1
        
    final = {s: (round(sums[s]/cnts[s], 2) if cnts[s]>0 else 0) for s in SCHEMAS_ORDER}
    return final, f"Succès ({limit} réponses)."

# --- 3. UTILS & SAVE ---
def clean_text(text):
    if not isinstance(text, str): return str(text)
    replacements = {"’": "'", "‘": "'", "“": '"', "”": '"', "–": "-", "…": "...", "œ": "oe", "«": '"', "»": '"', "€": "EUR"}
    for k,v in replacements.items(): text = text.replace(k,v)
    return text.encode('latin-1', 'replace').decode('latin-1')

def load_data():
    if os.path.exists(DB_FILE): return pd.read_csv(DB_FILE)
    return pd.DataFrame(columns=["Code_Couple", "Nom", "Date"] + SCHEMAS_ORDER)

def save_response(code, nom, scores):
    df = load_data()
    new_row = {"Code_Couple": code, "Nom": nom, "Date": datetime.now().strftime("%Y-%m-%d")}
    new_row.update(scores)
    df = df[~((df['Code_Couple'] == code) & (df['Nom'] == nom))]
    df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
    df.to_csv(DB_FILE, index=False)

def create_radar(d_A, d_B, n_A, n_B):
    cats = [SCHEMA_LIBRARY[s]['nom'] for s in SCHEMAS_ORDER]
    v_A = [d_A.get(s,0) for s in SCHEMAS_ORDER]; v_A += [v_A[0]]
    v_B = [d_B.get(s,0) for s in SCHEMAS_ORDER]; v_B += [v_B[0]]
    cats += [cats[0]]
    fig = go.Figure()
    fig.add_trace(go.Scatterpolar(r=v_A, theta=cats, fill='toself', name=n_A, line_color='#2980b9'))
    fig.add_trace(go.Scatterpolar(r=v_B, theta=cats, fill='toself', name=n_B, line_color='#e74c3c'))
    fig.update_layout(polar=dict(radialaxis=dict(visible=True, range=[0, 6])), template="plotly_white", margin=dict(t=30, b=30, l=40, r=40))
    return fig

# --- 4. PDF ---
class PDFExpert(FPDF):
    def header(self):
        self.set_fill_color(44, 62, 80)
        self.rect(0, 0, 210, 40, 'F')
        self.set_font('Arial', 'B', 24); self.set_text_color(255)
        self.set_xy(10, 10); self.cell(0, 15, clean_text("ALLIANCE & SCHEMAS"), 0, 1)
        self.set_font('Arial', 'I', 12)
        self.cell(0, 10, clean_text("Analyse Clinique, Théologique & Pastorale"), 0, 1)
        self.ln(15)
    def footer(self):
        self.set_y(-15); self.set_font('Arial', '', 8); self.set_text_color(128)
        self.cell(0, 10, clean_text(f"Page {self.page_no()}"), 0, 0, 'C')
    def draw_box(self, title, content, r, g, b):
        self.set_fill_color(r, g, b); self.set_font('Arial', 'B', 11); self.set_text_color(0)
        self.cell(0, 8, clean_text(f"  {title}"), 0, 1, 'L', 1)
        self.set_font('Arial', '', 10); self.set_text_color(50)
        self.multi_cell(0, 6, clean_text(content), border='L'); self.ln(3)

def generate_pdf(nA, dA, nB, dB, code):
    pdf = PDFExpert(); pdf.set_auto_page_break(True, 15); pdf.add_page()
    pdf.set_font('Arial', 'B', 16); pdf.set_text_color(0)
    pdf.cell(0, 10, clean_text(f"Dossier : {code}"), 0, 1)
    pdf.set_font('Arial', '', 12)
    pdf.cell(0, 8, clean_text(f"Partenaires : {nA} & {nB}"), 0, 1); pdf.ln(5)
    try:
        fig = create_radar(dA, dB, nA, nB)
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as t:
            fig.write_image(t.name, format="png", width=800, height=600, scale=2, engine="kaleido")
            pdf.image(t.name, x=15, y=70, w=180)
        pdf.set_y(230)
    except: pdf.cell(0, 10, "[Graphique manquant]", 0, 1)
    
    pdf.add_page(); pdf.set_font('Arial', 'B', 18); pdf.set_text_color(41, 128, 185)
    pdf.cell(0, 10, clean_text("ANALYSE DETAILLEE"), 0, 1); pdf.ln(5)
    
    def get_max(s): return max(dA.get(s,0), dB.get(s,0))
    ordered = sorted(SCHEMAS_ORDER, key=get_max, reverse=True)
    
    for s in ordered:
        mx = get_max(s)
        if mx >= 3.0:
            inf = SCHEMA_LIBRARY[s]
            pdf.ln(5); pdf.set_font('Arial', 'B', 14); pdf.set_text_color(192, 57, 43)
            pdf.cell(0, 10, clean_text(f"{inf['nom'].upper()} {'(!)' if mx>=5 else ''}"), 0, 1)
            pdf.set_font('Arial', 'B', 10); pdf.set_text_color(100)
            pdf.cell(0, 6, clean_text(f"A={dA.get(s,0)} | B={dB.get(s,0)}"), 0, 1); pdf.ln(2)
            
            pdf.draw_box("Clinique", inf['clinique'], 235, 245, 251)
            pdf.draw_box("Impact Couple", inf['couple'], 253, 237, 236)
            pdf.draw_box("Théologie (Cœur & Idoles)", inf['theologie'], 245, 245, 245)
            pdf.draw_box("Pastoral & Pratique", inf['pratique'], 233, 247, 239)
            
            pdf.set_font('Arial', 'I', 10); pdf.set_text_color(39, 174, 96)
            pdf.multi_cell(0, 6, clean_text(f"Verset : {inf['verset']}"))
            pdf.line(10, pdf.get_y()+2, 200, pdf.get_y()+2); pdf.ln(5)
    return pdf.output(dest='S').encode('latin-1', 'replace')

# --- 5. INTERFACE ---
st.sidebar.title("Navigation")
mode = st.sidebar.radio("Mode", ["🏠 Questionnaire", "💼 Espace Expert"])

if mode == "🏠 Questionnaire":
    st.title("Questionnaire Couple (YSQ-L3)")
    if 'usr' not in st.session_state: st.session_state.usr = None
    if 'pg' not in st.session_state: st.session_state.pg = 0
    if 'dat' not in st.session_state: st.session_state.dat = {}
    
    if not st.session_state.usr:
        with st.form("log"):
            c = st.text_input("Code"); n = st.text_input("Nom")
            if st.form_submit_button("Go"): st.session_state.usr={"c":c,"n":n}; st.rerun()
    else:
        # 232 questions paging
        PER_PAGE=40; tot=len(ALL_QUESTIONS); pgs=(tot//PER_PAGE)+1
        start=st.session_state.pg*PER_PAGE+1; end=min(start+PER_PAGE, tot+1)
        st.progress((st.session_state.pg+1)/pgs); st.caption(f"P {st.session_state.pg+1}/{pgs}")
        
        with st.form("q"):
            for i in range(start, end):
                if i in ALL_QUESTIONS:
                    st.markdown(f"**{i}. {ALL_QUESTIONS[i]['text']}**")
                    st.session_state.dat[i] = st.slider("", 1, 6, st.session_state.dat.get(i, 1), key=f"q{i}")
            if st.session_state.pg < pgs-1:
                if st.form_submit_button("Suivant"): st.session_state.pg+=1; st.rerun()
            else:
                if st.form_submit_button("Envoyer"):
                    m=get_schema_map_ordered(); sm={s:0 for s in SCHEMAS_ORDER}; cn={s:0 for s in SCHEMAS_ORDER}
                    for k,v in st.session_state.dat.items():
                        if k<=len(m): sc=m[k-1]; sm[sc]+=v; cn[sc]+=1
                    fin={s: (round(sm[s]/cn[s],2) if cn[s]>0 else 0) for s in SCHEMAS_ORDER}
                    save_response(st.session_state.usr['c'], st.session_state.usr['n'], fin)
                    st.success("Terminé !"); st.balloons()

elif mode == "💼 Espace Expert":
    st.title("Espace Expert")
    if st.sidebar.text_input("Password", type="password") == "Expert2024":
        
        with st.expander("📥 IMPORTER FICHIERS (.txt / .docx)", expanded=True):
            st.info("Vous pouvez uploader directement vos fichiers Word.")
            c1, c2 = st.columns(2)
            fA = c1.file_uploader("Fichier A", type=['txt','docx']); nA = c1.text_input("Nom A")
            fB = c2.file_uploader("Fichier B", type=['txt','docx']); nB = c2.text_input("Nom B")
            code = st.text_input("Code Dossier").strip().upper()
            
            if st.button("Importer"):
                if fA and fB and code:
                    txtA = extract_text_from_file(fA)
                    txtB = extract_text_from_file(fB)
                    sA, mA = parse_imported_text(txtA)
                    sB, mB = parse_imported_text(txtB)
                    
                    if sA and sB:
                        save_response(code, nA, sA); save_response(code, nB, sB)
                        st.success("Dossier créé !"); st.write(mA); st.write(mB)
                    else: st.error("Erreur lecture format [x/6].")
        
        st.divider()
        df = load_data()
        if not df.empty:
            sel = st.selectbox("Dossier", df['Code_Couple'].unique())
            sub = df[df['Code_Couple']==sel]
            if len(sub)>=2:
                rA=sub.iloc[0]; rB=sub.iloc[1]
                c1, c2 = st.columns([2,1])
                with c1: st.plotly_chart(create_radar(rA.to_dict(), rB.to_dict(), rA['Nom'], rB['Nom']))
                with c2:
                    pdf = generate_pdf(rA['Nom'], rA.to_dict(), rB['Nom'], rB.to_dict(), sel)
                    st.download_button("📥 Rapport PDF", pdf, f"Rap_{sel}.pdf")
